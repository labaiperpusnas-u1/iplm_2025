# -*- coding: utf-8 -*-
"""
IPLM 2025 — Final (Target Sampel 33.88% per Jenis) — TANPA Kinerja Relatif / Percentile
UPDATE (sesuai instruksi terbaru Anda) — TANPA mengubah pipeline lain:

FOKUS PEMBENAHAN (LLM + WORD):
1) Nilai Kepatuhan, Koleksi, Tenaga, Kinerja, Pelayanan, Pengelolaan:
   - TIDAK dikalikan 100.
   - Ditulis APA ADANYA dari kolom agregat aplikasi:
     Rata2_dim_kepatuhan, Rata2_sub_koleksi, Rata2_sub_sdm, Rata2_dim_kinerja,
     Rata2_sub_pelayanan, Rata2_sub_pengelolaan.
2) Nilai IPLM ditulis apa adanya: Indeks_Final_Wilayah_0_100.
3) LLM mengisi Interpretasi & Rekomendasi:
   - Interpretasi: deskriptif, kondisi riil berbasis relasi angka (lebih besar/kecil, gap, dominan, konsistensi),
     plus pemaknaan substantif dimensi (koleksi/sdm/pelayanan/pengelolaan) TANPA label normatif.
   - Rekomendasi: operasional, 2–3 butir ringkas, menaut ke pola angka (gap/ketimpangan/kontribusi).
4) Di bawah tabel Word: tambah deskripsi jumlah perpustakaan sumber data (dari tabel agregat wilayah × jenis / “gambar 2”):
   sekolah=..., umum=..., khusus=..., total=...

Catatan penting:
- Semua perhitungan dan dashboard tetap.
- Yang diubah hanya: (a) cara mengambil nilai untuk tabel Word (tanpa *100),
  (b) prompt LLM untuk isi interpretasi/rekomendasi agar nyambung dengan angka,
  (c) tambahan paragraf jumlah perpustakaan di bawah tabel Word.
"""

import os
import re
import time
import json
import math
import tempfile
from pathlib import Path

import gradio as gr
import numpy as np
import pandas as pd
import plotly.graph_objects as go
from sklearn.preprocessing import PowerTransformer

# python-docx (wajib kalau mau Word)
DOCX_AVAILABLE = True
try:
    from docx import Document
    from docx.shared import Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
except Exception:
    DOCX_AVAILABLE = False
    Document = None

# huggingface client (opsional)
HF_AVAILABLE = True
try:
    from huggingface_hub import InferenceClient
except Exception:
    HF_AVAILABLE = False
    InferenceClient = None


# ============================================================
# 1) KONFIGURASI
# ============================================================

DATA_FILE   = os.getenv("DATA_FILE",   "DATA CLEAN GABUNGAN SANGGAH-TIDAK SANGGAH ALL2026.xlsx")
POP_KAB     = os.getenv("POP_KAB",     "Data_populasi_Kab_kota_fixed.xlsx")
POP_PROV    = os.getenv("POP_PROV",    "Data_populasi_propinsi.xlsx")
POP_KHUSUS  = os.getenv("POP_KHUSUS",  "Data_populasi_perp_khusus.xlsx")

W_KEPATUHAN = float(os.getenv("W_KEPATUHAN", "0.30"))
W_KINERJA   = float(os.getenv("W_KINERJA",   "0.70"))

TARGET_RATIO = float(os.getenv("TARGET_RATIO", "0.3388"))

USE_LLM = True
LLM_MODEL_NAME = os.getenv("LLM_MODEL_NAME", "deepseek-ai/DeepSeek-R1-Distill-Qwen-32B")
HF_TOKEN = (
    os.getenv("HF_SECRET")
    or os.getenv("HF_TOKEN")
    or os.getenv("HUGGINGFACEHUB_API_TOKEN")
    or os.getenv("HF_API_TOKEN")
)
#meta-llama/Meta-Llama-3-8B-Instruct

# ============================================================
# 2) UTIL
# ============================================================

def _mtime(path_str: str):
    p = Path(path_str)
    return p.stat().st_mtime if p.exists() else None

def _canon(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(s).lower())

def _disp_text(x):
    if pd.isna(x):
        return None
    t = str(x).strip().upper()
    return " ".join(t.split())

def pick_col(df, candidates):
    if df is None or df.empty:
        return None
    for c in candidates:
        if c in df.columns:
            return c
    can_map = {_canon(c): c for c in df.columns}
    for c in candidates:
        k = _canon(c)
        if k in can_map:
            return can_map[k]
    return None

def coerce_num(val):
    if pd.isna(val):
        return np.nan
    t = str(val).strip()
    if t == "" or t in {"-", "–", "—", "NA", "N/A", "null", "NULL"}:
        return np.nan
    t = t.replace("\u00a0", " ").replace("Rp", "").replace("%", "")
    t = re.sub(r"[^0-9,.\-]", "", t)

    if t.count(".") > 1 and t.count(",") == 1:
        t = t.replace(".", "").replace(",", ".")
    elif t.count(",") > 1 and t.count(".") == 1:
        t = t.replace(",", "")
    elif t.count(",") == 1 and t.count(".") == 0:
        t = t.replace(",", ".")
    else:
        t = t.replace(",", "")

    try:
        return float(t)
    except Exception:
        return np.nan

def minmax_norm(s: pd.Series) -> pd.Series:
    x = pd.to_numeric(s, errors="coerce").astype(float)
    mn, mx = x.min(skipna=True), x.max(skipna=True)
    if pd.isna(mn) or pd.isna(mx) or mx == mn:
        return pd.Series(0.0, index=s.index)
    return (x - mn) / (mx - mn)

def norm_kew(v):
    if pd.isna(v):
        return None
    t = str(v).strip().upper()
    if "KAB" in t or "KOTA" in t:
        return "KAB/KOTA"
    if "PROV" in t:
        return "PROVINSI"
    if "PUSAT" in t or "NASIONAL" in t:
        return "PUSAT"
    return t

def norm_prov_disp(s):
    if pd.isna(s):
        return None
    t = str(s).strip().upper()
    t = t.replace("\u00a0", " ")
    t = " ".join(t.split())
    t = t.replace("PROPINSI", "PROVINSI")
    while t.startswith("PROVINSI PROVINSI "):
        t = t.replace("PROVINSI PROVINSI ", "PROVINSI ", 1)
    if t.startswith("PROVINSI "):
        name = t[len("PROVINSI "):].strip()
    else:
        name = t
    name = " ".join(name.split())
    if not name:
        return None
    return f"PROVINSI {name}"

def norm_prov_label(s):
    if pd.isna(s):
        return None
    t = str(s).strip().upper().replace("\u00a0", " ")
    t = " ".join(t.split())
    t = t.replace("PROPINSI", "PROVINSI")
    t = t.replace("PROVINSI", "").strip()
    return re.sub(r"[^A-Z0-9]+", "", t)

def norm_kab_label(s):
    if pd.isna(s):
        return None
    t = str(s).upper()
    t = t.replace("KABUPATEN", "KAB")
    t = t.replace("KAB.", "KAB")
    t = t.replace("KOTA ADMINISTRASI", "KOTA")
    t = t.replace("KOTA ADM.", "KOTA")
    t = t.replace("KOTA.", "KOTA")
    t = " ".join(t.split())
    return re.sub(r"[^A-Z0-9]+", "", t)

def safe_div(num, den):
    if den is None or pd.isna(den) or float(den) <= 0:
        return np.nan
    return float(num) / float(den)

def faktor_penyesuaian_total(n_total: float, target_total: float) -> float:
    if target_total is None or pd.isna(target_total) or float(target_total) <= 0:
        return 1.0
    if n_total is None or pd.isna(n_total) or float(n_total) < 0:
        n_total = 0.0
    return float(min(float(n_total) / float(target_total), 1.0))


# ============================================================
# 3) INDIKATOR IPLM
# ============================================================

koleksi_cols = [
    "JudulTercetak","EksemplarTercetak","JudulElektronik","EksemplarElektronik",
    "TambahJudulTercetak","TambahEksemplarTercetak",
    "TambahJudulElektronik","TambahEksemplarElektronik",
    "KomitmenAnggaranKoleksi"
]
sdm_cols = [
    "TenagaKualifikasiIlmuPerpustakaan",
    "TenagaFungsionalProfesional",
    "TenagaPKB",
    "AnggaranTenaga"
]
pelayanan_cols = [
    "PesertaBudayaBaca","PemustakaLuringDaring","PemustakaFasilitasTIK",
    "PemanfaatanJudulTercetak","PemanfaatanEksemplarTercetak",
    "PemanfaatanJudulElektronik","PemanfaatanEksemplarElektronik"
]
pengelolaan_cols = [
    "KegiatanBudayaBaca","KegiatanKerjasama","VariasiLayanan","Kebijakan","AnggaranLayanan"
]
all_indicators = koleksi_cols + sdm_cols + pelayanan_cols + pengelolaan_cols

alias_map_raw = {
    "j_judul_koleksi_tercetak": "JudulTercetak",
    "j_eksemplar_koleksi_tercetak": "EksemplarTercetak",
    "j_judul_koleksi_digital": "JudulElektronik",
    "j_eksemplar_koleksi_digital": "EksemplarElektronik",
    "tambah_judul_koleksi_tercetak": "TambahJudulTercetak",
    "tambah_eksemplar_koleksi_tercetak": "TambahEksemplarTercetak",
    "tambah_judul_koleksi_digital": "TambahJudulElektronik",
    "tambah_eksemplar_koleksi_digital": "TambahEksemplarElektronik",
    "j_anggaran_koleksi": "KomitmenAnggaranKoleksi",
    "j_tenaga_ilmu_perpus": "TenagaKualifikasiIlmuPerpustakaan",
    "j_tenaga_nonilmu_perpus": "TenagaFungsionalProfesional",
    "j_tenaga_pkb": "TenagaPKB",
    "j_anggaran_diklat_perpus": "AnggaranTenaga",
    "j_peserta_budaya_baca": "PesertaBudayaBaca",
    "j_pemustaka_luring_daring": "PemustakaLuringDaring",
    "j_pemustaka_fasilitas_tik": "PemustakaFasilitasTIK",
    "j_judul_koleksi_tercetak_termanfaat": "PemanfaatanJudulTercetak",
    "j_eksemplar_koleksi_tercetak_termanfaat": "PemanfaatanEksemplarTercetak",
    "j_judul_koleksi_digital_termanfaat": "PemanfaatanJudulElektronik",
    "j_eksemplar_koleksi_digital_termanfaat": "PemanfaatanEksemplarElektronik",
    "j_kegiatan_budaya_baca_peningkatan_literasi": "KegiatanBudayaBaca",
    "j_kerjasama_pengembangan_perpus": "KegiatanKerjasama",
    "j_variasi_layanan": "VariasiLayanan",
    "j_kebijakan_prosedur_pelayanan": "Kebijakan",
    "j_anggaran_peningkatan_pelayanan": "AnggaranLayanan",
}
alias_map = {_canon(k): v for k, v in alias_map_raw.items()}


# ============================================================
# 4) PIPELINE NASIONAL (LEVEL ENTITAS)
# ============================================================

def _mean_norm_cols(row, cols):
    vals = []
    for c in cols:
        k = f"norm_{c}"
        if k in row.index:
            v = row[k]
            if pd.isna(v):
                v = 0.0
            vals.append(float(v))
    return float(np.mean(vals)) if vals else 0.0

def prepare_global(df_src: pd.DataFrame) -> pd.DataFrame:
    if df_src is None or df_src.empty:
        return df_src

    df = df_src.copy()

    rename_map = {}
    for col in df.columns:
        c = _canon(col)
        if c in alias_map:
            rename_map[col] = alias_map[c]
        else:
            for tgt in all_indicators:
                if c == _canon(tgt):
                    rename_map[col] = tgt
                    break
    if rename_map:
        df = df.rename(columns=rename_map)

    available = [c for c in all_indicators if c in df.columns]
    for c in available:
        df[c] = df[c].apply(coerce_num)

    for c in available:
        x = pd.to_numeric(df[c], errors="coerce").astype(float).values
        mask = ~np.isnan(x)
        transformed = np.full_like(x, np.nan, dtype=float)
        if mask.sum() > 1:
            pt = PowerTransformer(method="yeo-johnson", standardize=False)
            transformed[mask] = pt.fit_transform(x[mask].reshape(-1, 1)).ravel()
        else:
            transformed[mask] = x[mask]
        df[f"norm_{c}"] = minmax_norm(pd.Series(transformed, index=df.index))

    df["sub_koleksi"]     = df.apply(lambda r: _mean_norm_cols(r, [c for c in koleksi_cols if c in available]), axis=1)
    df["sub_sdm"]         = df.apply(lambda r: _mean_norm_cols(r, [c for c in sdm_cols if c in available]), axis=1)
    df["sub_pelayanan"]   = df.apply(lambda r: _mean_norm_cols(r, [c for c in pelayanan_cols if c in available]), axis=1)
    df["sub_pengelolaan"] = df.apply(lambda r: _mean_norm_cols(r, [c for c in pengelolaan_cols if c in available]), axis=1)

    df["dim_kepatuhan"] = df[["sub_koleksi","sub_sdm"]].mean(axis=1)
    df["dim_kinerja"]   = df[["sub_pelayanan","sub_pengelolaan"]].mean(axis=1)

    df["Indeks_Dasar_0_100"] = 100 * (W_KEPATUHAN * df["dim_kepatuhan"] + W_KINERJA * df["dim_kinerja"])

    for c in ["sub_koleksi","sub_sdm","sub_pelayanan","sub_pengelolaan","dim_kepatuhan","dim_kinerja","Indeks_Dasar_0_100"]:
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0.0)

    return df


# ============================================================
# 5) CACHE LOADER (NO UPLOAD)
# ============================================================

_CACHE = {"key": None, "df_all": None, "df_raw": None, "pop_kab": None, "pop_prov": None, "pop_khusus": None, "meta": None, "info": None}

def _parse_pop_khusus(path_xlsx: str) -> pd.DataFrame:
    df = pd.read_excel(path_xlsx)
    if df is None or df.empty:
        return pd.DataFrame()

    c_mix = pick_col(df, [
        "Propinsi/Kab/kota", "Propinsi/Kab/Kota", "Provinsi/Kab/Kota",
        "Provinsi/Kab/kota", "Provinsi/Kabupaten/Kota",
        "Wilayah", "Nama Wilayah"
    ])
    if c_mix is None:
        raise ValueError("POP_KHUSUS: kolom gabungan Provinsi/Kab/Kota tidak ditemukan.")

    c_pop = pick_col(df, ["POP_KHUSUS", "Pop_Khusus", "pop_khusus"])
    if c_pop is None:
        raise ValueError("POP_KHUSUS: kolom 'POP_KHUSUS' tidak ditemukan.")

    mix = df[c_mix].astype(str).fillna("").str.strip()
    pop_series = df[c_pop].apply(coerce_num)

    rows = []
    current_prov = None

    for m, pval in zip(mix.tolist(), pop_series.tolist()):
        mm = _disp_text(m) or ""
        if mm == "":
            continue

        if mm.startswith("PROVINSI "):
            prov_name = mm.replace("PROVINSI", "").strip()
            current_prov = prov_name
            rows.append({
                "LEVEL": "PROV",
                "Provinsi_Label": f"PROVINSI {prov_name}",
                "Kab_Kota_Label": None,
                "Pop_Total_Jenis": pval,
            })
            continue

        rows.append({
            "LEVEL": "KAB",
            "Provinsi_Label": f"PROVINSI {current_prov}" if current_prov else None,
            "Kab_Kota_Label": mm,
            "Pop_Total_Jenis": pval,
        })

    pop = pd.DataFrame(rows)
    if pop.empty:
        return pop

    pop["Pop_Total_Jenis"] = pd.to_numeric(pop["Pop_Total_Jenis"], errors="coerce").fillna(0.0)
    pop["prov_key"] = pop["Provinsi_Label"].apply(norm_prov_label)
    pop["kab_key"] = pop["Kab_Kota_Label"].apply(norm_kab_label) if "Kab_Kota_Label" in pop.columns else None
    return pop

def load_default_files(force=False):
    key = (
        DATA_FILE, POP_KAB, POP_PROV, POP_KHUSUS,
        _mtime(DATA_FILE), _mtime(POP_KAB), _mtime(POP_PROV), _mtime(POP_KHUSUS)
    )

    if (not force) and _CACHE["key"] == key and _CACHE["df_all"] is not None:
        return _CACHE["df_all"], _CACHE["df_raw"], _CACHE["pop_kab"], _CACHE["pop_prov"], _CACHE["pop_khusus"], _CACHE["meta"], _CACHE["info"]

    for p, label in [(DATA_FILE, "DM"), (POP_KAB, "POP_KAB"), (POP_PROV, "POP_PROV"), (POP_KHUSUS, "POP_KHUSUS")]:
        if not Path(p).exists():
            info = f"File tidak ditemukan ({label}): {p}"
            _CACHE.update({"key": key, "df_all": None, "df_raw": None, "pop_kab": None, "pop_prov": None, "pop_khusus": None, "meta": {}, "info": info})
            return None, None, None, None, None, {}, info

    fp = Path(DATA_FILE)
    xls = pd.ExcelFile(fp)
    frames = [pd.read_excel(fp, sheet_name=s) for s in xls.sheet_names]
    df_raw = pd.concat(frames, ignore_index=True, sort=False)

    prov_col  = pick_col(df_raw, ["provinsi", "Provinsi", "PROVINSI"])
    kab_col   = pick_col(df_raw, ["kab/kota", "Kab/Kota", "Kab_Kota", "KAB/KOTA", "kabupaten_kota", "Kabupaten/Kota", "kabupaten kota", "kota", "kab_kota"])
    kew_col   = pick_col(df_raw, ["kewenangan", "jenis_kewenangan", "Kewenangan", "KEWENANGAN"])
    jenis_col = pick_col(df_raw, ["jenis_perpustakaan", "Jenis Perpustakaan", "JENIS_PERPUSTAKAAN"])
    nama_col  = pick_col(df_raw, ["nm_perpustakaan","nama_perpustakaan","Nama Perpustakaan","nm_instansi_lembaga","nm_perpus"])

    missing = []
    if prov_col is None:  missing.append("Provinsi")
    if kab_col is None:   missing.append("Kab/Kota")
    if kew_col is None:   missing.append("Kewenangan")
    if jenis_col is None: missing.append("Jenis Perpustakaan")
    if missing:
        info = f"Kolom wajib tidak ditemukan di DM: {', '.join(missing)}"
        _CACHE.update({"key": key, "df_all": None, "df_raw": None, "pop_kab": None, "pop_prov": None, "pop_khusus": None, "meta": {}, "info": info})
        return None, None, None, None, None, {}, info

    val_map_jenis = {
        "PERPUSTAKAAN SEKOLAH": "sekolah", "SEKOLAH": "sekolah",
        "PERPUSTAKAAN UMUM": "umum", "UMUM": "umum", "PERPUSTAKAAN DAERAH": "umum",
        "PERPUSTAKAAN KHUSUS": "khusus", "KHUSUS": "khusus",
    }

    df_raw["KEW_NORM"]  = df_raw[kew_col].apply(norm_kew)
    df_raw["_dataset"]  = df_raw[jenis_col].astype(str).str.strip().str.upper().map(val_map_jenis)
    df_raw["PROV_DISP"] = df_raw[prov_col].apply(norm_prov_disp)
    df_raw["KAB_DISP"]  = df_raw[kab_col].apply(_disp_text)
    df_raw["prov_key"]  = df_raw["PROV_DISP"].apply(norm_prov_label)
    df_raw["kab_key"]   = df_raw["KAB_DISP"].apply(norm_kab_label)

    if nama_col and nama_col in df_raw.columns:
        kcols = [prov_col, kab_col, kew_col, jenis_col, nama_col]
    else:
        kcols = [prov_col, kab_col, kew_col, jenis_col]

    tmp = df_raw[kcols].astype(str).fillna("").apply(lambda s: s.str.strip(), axis=0)
    df_raw["_row_key"] = tmp.apply(lambda r: "||".join(r.values.tolist()), axis=1).apply(_canon)
    before = len(df_raw)
    df_raw = df_raw.drop_duplicates(subset=["_row_key"], keep="first").copy()
    after = len(df_raw)

    # POP KAB
    pk = pd.read_excel(POP_KAB)
    c_kab  = pick_col(pk, ["KABUPATEN_KOTA","Kab/Kota","Kabupaten/Kota","KAB/KOTA","Kabupaten_Kota","kab_kota","kabupaten_kota"])
    c_prov = pick_col(pk, ["PROVINSI","Provinsi","provinsi"])
    if c_kab is None:
        info = "POP_KAB: wajib ada kolom Kab/Kota."
        _CACHE.update({"key": key, "df_all": None, "df_raw": None, "pop_kab": None, "pop_prov": None, "pop_khusus": None, "meta": {}, "info": info})
        return None, None, None, None, None, {}, info

    pop_kab = pk.copy()
    pop_kab["Kab_Kota_Label"] = pk[c_kab].astype(str).str.strip()
    pop_kab["Provinsi_Label"] = pk[c_prov].astype(str).str.strip() if c_prov else ""
    pop_kab["kab_key"] = pop_kab["Kab_Kota_Label"].apply(norm_kab_label)
    pop_kab = pop_kab.groupby("kab_key", as_index=False).first()

    # POP PROV
    pp = pd.read_excel(POP_PROV)
    c_pr = pick_col(pp, ["Provinsi","PROVINSI","provinsi","Propinsi","PROPINSI","propinsi"])
    if c_pr is None:
        info = "POP_PROV: wajib ada kolom Provinsi."
        _CACHE.update({"key": key, "df_all": None, "df_raw": None, "pop_kab": None, "pop_prov": None, "pop_khusus": None, "meta": {}, "info": info})
        return None, None, None, None, None, {}, info

    pop_prov = pp.copy()
    pop_prov["Provinsi_Label"] = pp[c_pr].astype(str).str.strip()
    pop_prov["prov_key"] = pop_prov["Provinsi_Label"].apply(norm_prov_label)
    pop_prov = pop_prov.groupby("prov_key", as_index=False).first()

    # POP KHUSUS
    try:
        pop_khusus = _parse_pop_khusus(POP_KHUSUS)
    except Exception as e:
        info = f"POP_KHUSUS gagal dibaca: {repr(e)}"
        _CACHE.update({"key": key, "df_all": None, "df_raw": None, "pop_kab": None, "pop_prov": None, "pop_khusus": None, "meta": {}, "info": info})
        return None, None, None, None, None, {}, info

    df_all = prepare_global(df_raw)
    meta = dict(prov_col=prov_col, kab_col=kab_col, kew_col=kew_col, jenis_col=jenis_col, nama_col=nama_col)

    info = (
        f"Mode NO UPLOAD (cache aktif)\n"
        f"DM: {fp.name} | Baris: {before} -> dedup: {after}\n"
        f"POP_KAB: {Path(POP_KAB).name} (n={len(pop_kab)})\n"
        f"POP_PROV: {Path(POP_PROV).name} (n={len(pop_prov)})\n"
        f"POP_KHUSUS: {Path(POP_KHUSUS).name} (n={len(pop_khusus)})\n"
        f"TARGET sampel per jenis: {TARGET_RATIO*100:.2f}%\n"
        f"mtime: DM={time.ctime(_mtime(DATA_FILE))} | Kab={time.ctime(_mtime(POP_KAB))} | Prov={time.ctime(_mtime(POP_PROV))} | Khusus={time.ctime(_mtime(POP_KHUSUS))}"
    )

    _CACHE.update({
        "key": key,
        "df_all": df_all,
        "df_raw": df_raw,
        "pop_kab": pop_kab,
        "pop_prov": pop_prov,
        "pop_khusus": pop_khusus,
        "meta": meta,
        "info": info
    })
    return df_all, df_raw, pop_kab, pop_prov, pop_khusus, meta, info


# ============================================================
# 6) FAKTOR WILAYAH — PER JENIS (TARGET 33.88%)
# ============================================================

def build_faktor_wilayah_jenis(df_filtered, pop_kab, pop_prov, pop_khusus, kew_value):
    if df_filtered is None or df_filtered.empty:
        return pd.DataFrame()

    kew_norm = str(kew_value or "").upper()
    df = df_filtered.copy()
    df = df[df["_dataset"].isin(["sekolah", "umum", "khusus"])].copy()
    if df.empty:
        return pd.DataFrame()

    jenis_list = ["sekolah", "umum", "khusus"]

    if "PROV" in kew_norm:
        key_col, label_col, label_name, mode = "prov_key", "PROV_DISP", "Provinsi", "PROV"
        base_pop = pop_prov.copy() if (pop_prov is not None and not pop_prov.empty) else pd.DataFrame()
        if not base_pop.empty and "prov_key" not in base_pop.columns:
            base_pop["prov_key"] = base_pop["Provinsi_Label"].apply(norm_prov_label)
        base_pop = base_pop.set_index("prov_key") if (not base_pop.empty and "prov_key" in base_pop.columns) else pd.DataFrame().set_index(pd.Index([]))
    else:
        key_col, label_col, label_name, mode = "kab_key", "KAB_DISP", "Kab/Kota", "KAB"
        base_pop = pop_kab.copy() if (pop_kab is not None and not pop_kab.empty) else pd.DataFrame()
        if not base_pop.empty and "kab_key" not in base_pop.columns:
            base_pop["kab_key"] = base_pop["Kab_Kota_Label"].apply(norm_kab_label)
        base_pop = base_pop.set_index("kab_key") if (not base_pop.empty and "kab_key" in base_pop.columns) else pd.DataFrame().set_index(pd.Index([]))

    base_keys = df[[key_col, label_col]].drop_duplicates().rename(columns={key_col: "group_key", label_col: label_name})
    full = base_keys.assign(_tmp=1).merge(pd.DataFrame({"Jenis": jenis_list, "_tmp": 1}), on="_tmp").drop(columns="_tmp")

    cnt = (
        df.groupby([key_col, label_col, "_dataset"], dropna=False)
          .size()
          .reset_index(name="n_jenis")
          .rename(columns={key_col: "group_key", label_col: label_name, "_dataset": "Jenis"})
    )
    cnt["Jenis"] = cnt["Jenis"].astype(str).str.lower().str.strip()

    base_n = full.merge(cnt, on=["group_key", label_name, "Jenis"], how="left")
    base_n["n_jenis"] = pd.to_numeric(base_n["n_jenis"], errors="coerce").fillna(0).astype(int)

    base_n["target_total_33_88_jenis"] = 0.0
    base_n["pop_total_jenis"] = 0.0

    if not base_pop.empty:
        if mode == "KAB":
            pop_sekolah = pd.to_numeric(base_pop.get("jumlah_populasi_sekolah_base", base_pop.get("jumlah_populasi_sekolah", 0)), errors="coerce").fillna(0.0)
            pop_umum    = pd.to_numeric(base_pop.get("jumlah_populasi_umum_base", base_pop.get("jumlah_populasi_umum", 0)), errors="coerce").fillna(0.0)
            tgt_sekolah = pop_sekolah * float(TARGET_RATIO)
            tgt_umum    = pop_umum * float(TARGET_RATIO)
        else:
            sma = pd.to_numeric(base_pop.get("sma ", base_pop.get("sma", 0)), errors="coerce").fillna(0.0)
            smk = pd.to_numeric(base_pop.get("smk", 0), errors="coerce").fillna(0.0)
            slb = pd.to_numeric(base_pop.get("slb", 0), errors="coerce").fillna(0.0)
            pop_sekolah = sma + smk + slb
            tgt_sekolah = pop_sekolah * float(TARGET_RATIO)
            pop_umum = pd.to_numeric(base_pop.get("perpus_umum_prop", 0), errors="coerce").fillna(0.0)
            tgt_umum = pop_umum * float(TARGET_RATIO)

        m = base_n["Jenis"].eq("sekolah")
        base_n.loc[m, "pop_total_jenis"] = base_n.loc[m, "group_key"].map(pop_sekolah).fillna(0.0).values
        base_n.loc[m, "target_total_33_88_jenis"] = base_n.loc[m, "group_key"].map(tgt_sekolah).fillna(0.0).values

        m = base_n["Jenis"].eq("umum")
        base_n.loc[m, "pop_total_jenis"] = base_n.loc[m, "group_key"].map(pop_umum).fillna(0.0).values
        base_n.loc[m, "target_total_33_88_jenis"] = base_n.loc[m, "group_key"].map(tgt_umum).fillna(0.0).values

    if pop_khusus is not None and not pop_khusus.empty:
        pk = pop_khusus.copy()
        pk["Pop_Total_Jenis"] = pd.to_numeric(pk.get("Pop_Total_Jenis", 0), errors="coerce").fillna(0.0)

        if mode == "PROV":
            pk_prov = pk[pk["LEVEL"].astype(str).str.upper() == "PROV"].copy()
            pk_map = pk_prov.groupby("prov_key", as_index=True).agg(pop=("Pop_Total_Jenis", "sum"))
            pop_series = pk_map["pop"]
        else:
            pk_kab = pk[pk["LEVEL"].astype(str).str.upper() == "KAB"].copy()
            pk_map = pk_kab.groupby("kab_key", as_index=True).agg(pop=("Pop_Total_Jenis", "sum"))
            pop_series = pk_map["pop"]

        tgt_series = pop_series * float(TARGET_RATIO)

        m = base_n["Jenis"].eq("khusus")
        base_n.loc[m, "pop_total_jenis"] = base_n.loc[m, "group_key"].map(pop_series).fillna(0.0).values
        base_n.loc[m, "target_total_33_88_jenis"] = base_n.loc[m, "group_key"].map(tgt_series).fillna(0.0).values

    base_n["target_total_33_88_jenis"] = pd.to_numeric(base_n["target_total_33_88_jenis"], errors="coerce").fillna(0.0)
    base_n["pop_total_jenis"] = pd.to_numeric(base_n["pop_total_jenis"], errors="coerce").fillna(0.0)

    m_need_pop = (base_n["pop_total_jenis"] <= 0) & (base_n["target_total_33_88_jenis"] > 0)
    base_n.loc[m_need_pop, "pop_total_jenis"] = base_n.loc[m_need_pop, "target_total_33_88_jenis"] / float(TARGET_RATIO)

    base_n["faktor_penyesuaian_jenis"] = [
        faktor_penyesuaian_total(n, t)
        for n, t in zip(
            pd.to_numeric(base_n["n_jenis"], errors="coerce").fillna(0).astype(float),
            pd.to_numeric(base_n["target_total_33_88_jenis"], errors="coerce").fillna(0).astype(float),
        )
    ]

    base_n["coverage_jenis_%"] = [
        (safe_div(n, p) * 100.0) if (p is not None and not pd.isna(p) and float(p) > 0) else 0.0
        for n, p in zip(
            pd.to_numeric(base_n["n_jenis"], errors="coerce").fillna(0).astype(float),
            pd.to_numeric(base_n["pop_total_jenis"], errors="coerce").fillna(0).astype(float),
        )
    ]

    base_n["gap_target33_88_jenis"] = [
        max(float(t) - float(n), 0.0)
        for n, t in zip(
            pd.to_numeric(base_n["n_jenis"], errors="coerce").fillna(0).astype(float),
            pd.to_numeric(base_n["target_total_33_88_jenis"], errors="coerce").fillna(0).astype(float),
        )
    ]

    base_n["target_total_33_88_jenis"] = pd.to_numeric(base_n["target_total_33_88_jenis"], errors="coerce").fillna(0).round(0).astype(int)
    base_n["pop_total_jenis"]          = pd.to_numeric(base_n["pop_total_jenis"], errors="coerce").fillna(0).round(0).astype(int)
    base_n["coverage_jenis_%"]         = pd.to_numeric(base_n["coverage_jenis_%"], errors="coerce").fillna(0.0).round(2)
    base_n["faktor_penyesuaian_jenis"] = pd.to_numeric(base_n["faktor_penyesuaian_jenis"], errors="coerce").fillna(1.0).round(3)
    base_n["gap_target33_88_jenis"]    = pd.to_numeric(base_n["gap_target33_88_jenis"], errors="coerce").fillna(0).round(0).astype(int)

    return base_n


# ============================================================
# 7) AGREGAT WILAYAH × JENIS
# ============================================================

def build_agg_wilayah_jenis(df_filtered, faktor_wilayah_jenis, kew_value):
    if df_filtered is None or df_filtered.empty:
        return pd.DataFrame()

    kew_norm = str(kew_value or "").upper()
    df = df_filtered.copy()

    if "PROV" in kew_norm:
        key_col, label_col, label_name = "prov_key", "PROV_DISP", "Provinsi"
    else:
        key_col, label_col, label_name = "kab_key", "KAB_DISP", "Kab/Kota"

    df = df[df["_dataset"].isin(["sekolah", "umum", "khusus"])].copy()
    if df.empty:
        return pd.DataFrame()

    jenis_list = ["sekolah", "umum", "khusus"]

    base_keys = df[[key_col, label_col]].drop_duplicates().rename(columns={key_col: "group_key", label_col: label_name})
    full = base_keys.assign(_tmp=1).merge(pd.DataFrame({"Jenis": jenis_list, "_tmp": 1}), on="_tmp").drop(columns="_tmp")

    agg_real = df.groupby([key_col, label_col, "_dataset"], dropna=False).agg(
        Jumlah=("Indeks_Dasar_0_100", "size"),
        Rata2_sub_koleksi=("sub_koleksi", "mean"),
        Rata2_sub_sdm=("sub_sdm", "mean"),
        Rata2_sub_pelayanan=("sub_pelayanan", "mean"),
        Rata2_sub_pengelolaan=("sub_pengelolaan", "mean"),
        Rata2_dim_kepatuhan=("dim_kepatuhan", "mean"),
        Rata2_dim_kinerja=("dim_kinerja", "mean"),
        Indeks_Dasar_Agregat_0_100=("Indeks_Dasar_0_100", "mean"),
    ).reset_index().rename(columns={key_col: "group_key", label_col: label_name, "_dataset": "Jenis"})

    agg_real["Jenis"] = agg_real["Jenis"].astype(str).str.lower().str.strip()

    agg = full.merge(agg_real, on=["group_key", label_name, "Jenis"], how="left")
    for c in ["Jumlah","Rata2_sub_koleksi","Rata2_sub_sdm","Rata2_sub_pelayanan","Rata2_sub_pengelolaan",
              "Rata2_dim_kepatuhan","Rata2_dim_kinerja","Indeks_Dasar_Agregat_0_100"]:
        if c in agg.columns:
            agg[c] = pd.to_numeric(agg[c], errors="coerce").fillna(0.0)
    agg["Jumlah"] = agg["Jumlah"].round(0).astype(int)

    if faktor_wilayah_jenis is None or faktor_wilayah_jenis.empty:
        agg["faktor_penyesuaian_jenis"] = 1.0
    else:
        fw = faktor_wilayah_jenis.copy()
        fw["Jenis"] = fw["Jenis"].astype(str).str.lower().str.strip()
        keep = ["group_key", label_name, "Jenis",
                "faktor_penyesuaian_jenis", "target_total_33_88_jenis", "pop_total_jenis",
                "coverage_jenis_%", "gap_target33_88_jenis", "n_jenis"]
        fw = fw[[c for c in keep if c in fw.columns]].copy()
        agg = agg.merge(fw, on=["group_key", label_name, "Jenis"], how="left")
        agg["faktor_penyesuaian_jenis"] = pd.to_numeric(agg["faktor_penyesuaian_jenis"], errors="coerce").fillna(1.0)

    agg["Indeks_Final_Agregat_0_100"] = (
        pd.to_numeric(agg["Indeks_Dasar_Agregat_0_100"], errors="coerce").fillna(0.0)
        * pd.to_numeric(agg["faktor_penyesuaian_jenis"], errors="coerce").fillna(1.0)
    )

    for c in [
        "Rata2_sub_koleksi","Rata2_sub_sdm","Rata2_sub_pelayanan","Rata2_sub_pengelolaan",
        "Rata2_dim_kepatuhan","Rata2_dim_kinerja"
    ]:
        if c in agg.columns:
            agg[c] = pd.to_numeric(agg[c], errors="coerce").fillna(0.0).round(3)
    for c in ["Indeks_Dasar_Agregat_0_100","Indeks_Final_Agregat_0_100"]:
        if c in agg.columns:
            agg[c] = pd.to_numeric(agg[c], errors="coerce").fillna(0.0).round(2)

    agg["faktor_penyesuaian_jenis"] = pd.to_numeric(agg["faktor_penyesuaian_jenis"], errors="coerce").fillna(1.0).round(3)
    return agg


# ============================================================
# 8) AGREGAT WILAYAH (KESELURUHAN) — avg3 dari 3 jenis
# ============================================================

def build_agg_wilayah_total_from_jenis(agg_jenis, faktor_wilayah_jenis, kew_value):
    if agg_jenis is None or agg_jenis.empty:
        return pd.DataFrame()

    kew_norm = str(kew_value or "").upper()
    label_name = "Provinsi" if "PROV" in kew_norm else "Kab/Kota"
    jenis_list = ["sekolah", "umum", "khusus"]

    a = agg_jenis.copy()
    a["Jenis"] = a["Jenis"].astype(str).str.lower().str.strip()

    base_keys = a[["group_key", label_name]].drop_duplicates()
    full = base_keys.assign(_tmp=1).merge(pd.DataFrame({"Jenis": jenis_list, "_tmp": 1}), on="_tmp").drop(columns="_tmp")

    cols_present = [c for c in [
        "Jumlah",
        "Rata2_sub_koleksi","Rata2_sub_sdm","Rata2_sub_pelayanan","Rata2_sub_pengelolaan",
        "Rata2_dim_kepatuhan","Rata2_dim_kinerja",
        "Indeks_Dasar_Agregat_0_100",
        "Indeks_Final_Agregat_0_100",
    ] if c in a.columns]

    full = full.merge(a[["group_key", label_name, "Jenis"] + cols_present],
                      on=["group_key", label_name, "Jenis"], how="left")
    for c in cols_present:
        full[c] = pd.to_numeric(full[c], errors="coerce").fillna(0.0)

    out = full.groupby(["group_key", label_name], as_index=False).agg(
        n_total=("Jumlah", "sum"),
        Rata2_sub_koleksi=("Rata2_sub_koleksi", "mean"),
        Rata2_sub_sdm=("Rata2_sub_sdm", "mean"),
        Rata2_sub_pelayanan=("Rata2_sub_pelayanan", "mean"),
        Rata2_sub_pengelolaan=("Rata2_sub_pengelolaan", "mean"),
        Rata2_dim_kepatuhan=("Rata2_dim_kepatuhan", "mean"),
        Rata2_dim_kinerja=("Rata2_dim_kinerja", "mean"),
        Indeks_Dasar_Agregat_0_100=("Indeks_Dasar_Agregat_0_100", "mean"),
        Indeks_Final_Wilayah_0_100=("Indeks_Final_Agregat_0_100", "mean"),
    )

    for c in ["Rata2_sub_koleksi","Rata2_sub_sdm","Rata2_sub_pelayanan","Rata2_sub_pengelolaan","Rata2_dim_kepatuhan","Rata2_dim_kinerja"]:
        out[c] = pd.to_numeric(out[c], errors="coerce").fillna(0.0).round(3)
    for c in ["Indeks_Dasar_Agregat_0_100","Indeks_Final_Wilayah_0_100"]:
        out[c] = pd.to_numeric(out[c], errors="coerce").fillna(0.0).round(2)
    out["n_total"] = pd.to_numeric(out["n_total"], errors="coerce").fillna(0).round(0).astype(int)

    return out


# ============================================================
# 9) SUMMARY (PER JENIS) + KESELURUHAN
# ============================================================

def build_summary_per_jenis(agg_jenis, agg_total):
    jenis_list = ["sekolah", "umum", "khusus"]

    def _row_default(jenis):
        return {
            "Jenis": jenis,
            "Jumlah_Wilayah": 0,
            "Total_Perpus": 0,
            "Pop_Total_Jenis": 0,
            "Target33_88_Total_Jenis": 0,
            "Terkumpul_Jenis": 0,
            "Coverage_Target33_88_Jenis_%": 0.0,
            "Indeks_Dasar_0_100": 0.0,
            "Indeks_Final_Disesuaikan_0_100": 0.0,
            "Penyesuaian_Poin": 0.0,
        }

    rows_by_jenis = {j: _row_default(j) for j in jenis_list}

    if agg_jenis is not None and not agg_jenis.empty:
        a = agg_jenis.copy()
        a["Jenis"] = a["Jenis"].astype(str).str.lower().str.strip()
        for c in ["Jumlah","Indeks_Dasar_Agregat_0_100","Indeks_Final_Agregat_0_100","pop_total_jenis","target_total_33_88_jenis"]:
            if c in a.columns:
                a[c] = pd.to_numeric(a[c], errors="coerce").fillna(0)

        for jenis in jenis_list:
            sub = a[a["Jenis"] == jenis].copy()
            if sub.empty:
                continue

            jumlah_wilayah = int(sub.shape[0])
            terkumpul = int(pd.to_numeric(sub.get("Jumlah", 0), errors="coerce").fillna(0).sum())
            pop_total = int(pd.to_numeric(sub.get("pop_total_jenis", 0), errors="coerce").fillna(0).sum())
            target3388 = int(pd.to_numeric(sub.get("target_total_33_88_jenis", 0), errors="coerce").fillna(0).sum())

            coverage = (terkumpul / target3388 * 100.0) if target3388 > 0 else 0.0
            dasar = float(pd.to_numeric(sub.get("Indeks_Dasar_Agregat_0_100", 0), errors="coerce").fillna(0).mean())
            final = float(pd.to_numeric(sub.get("Indeks_Final_Agregat_0_100", 0), errors="coerce").fillna(0).mean())

            rows_by_jenis[jenis] = {
                "Jenis": jenis,
                "Jumlah_Wilayah": jumlah_wilayah,
                "Total_Perpus": terkumpul,
                "Pop_Total_Jenis": pop_total,
                "Target33_88_Total_Jenis": target3388,
                "Terkumpul_Jenis": terkumpul,
                "Coverage_Target33_88_Jenis_%": float(coverage),
                "Indeks_Dasar_0_100": float(dasar),
                "Indeks_Final_Disesuaikan_0_100": float(final),
                "Penyesuaian_Poin": float(final - dasar),
            }

    rows = [rows_by_jenis[j] for j in jenis_list]

    dasar_all = (rows_by_jenis["sekolah"]["Indeks_Dasar_0_100"]
                 + rows_by_jenis["umum"]["Indeks_Dasar_0_100"]
                 + rows_by_jenis["khusus"]["Indeks_Dasar_0_100"]) / 3.0

    final_all = (rows_by_jenis["sekolah"]["Indeks_Final_Disesuaikan_0_100"]
                 + rows_by_jenis["umum"]["Indeks_Final_Disesuaikan_0_100"]
                 + rows_by_jenis["khusus"]["Indeks_Final_Disesuaikan_0_100"]) / 3.0

    pop_all = int(rows_by_jenis["sekolah"]["Pop_Total_Jenis"]
                  + rows_by_jenis["umum"]["Pop_Total_Jenis"]
                  + rows_by_jenis["khusus"]["Pop_Total_Jenis"])

    target_all = int(rows_by_jenis["sekolah"]["Target33_88_Total_Jenis"]
                     + rows_by_jenis["umum"]["Target33_88_Total_Jenis"]
                     + rows_by_jenis["khusus"]["Target33_88_Total_Jenis"])

    terkumpul_all = int(rows_by_jenis["sekolah"]["Terkumpul_Jenis"]
                        + rows_by_jenis["umum"]["Terkumpul_Jenis"]
                        + rows_by_jenis["khusus"]["Terkumpul_Jenis"])

    coverage_all = (terkumpul_all / target_all * 100.0) if target_all > 0 else 0.0

    jumlah_wilayah_all = int(agg_total.shape[0]) if (agg_total is not None and not agg_total.empty) else int(
        max(rows_by_jenis["sekolah"]["Jumlah_Wilayah"],
            rows_by_jenis["umum"]["Jumlah_Wilayah"],
            rows_by_jenis["khusus"]["Jumlah_Wilayah"])
    )

    rows.append({
        "Jenis": "keseluruhan",
        "Jumlah_Wilayah": jumlah_wilayah_all,
        "Total_Perpus": terkumpul_all,
        "Pop_Total_Jenis": pop_all,
        "Target33_88_Total_Jenis": target_all,
        "Terkumpul_Jenis": terkumpul_all,
        "Coverage_Target33_88_Jenis_%": float(coverage_all),
        "Indeks_Dasar_0_100": float(dasar_all),
        "Indeks_Final_Disesuaikan_0_100": float(final_all),
        "Penyesuaian_Poin": float(final_all - dasar_all),
    })

    out = pd.DataFrame(rows)
    for c in ["Jumlah_Wilayah","Total_Perpus","Pop_Total_Jenis","Target33_88_Total_Jenis","Terkumpul_Jenis"]:
        out[c] = pd.to_numeric(out[c], errors="coerce").fillna(0).round(0).astype(int)
    for c in ["Coverage_Target33_88_Jenis_%","Indeks_Dasar_0_100","Indeks_Final_Disesuaikan_0_100","Penyesuaian_Poin"]:
        out[c] = pd.to_numeric(out[c], errors="coerce").fillna(0.0).round(2)
    return out


# ============================================================
# 10) DETAIL ENTITAS (Final menempel dari wilayah)
# ============================================================

def attach_final_to_detail(df_filtered, agg_total, meta, kew_value):
    if df_filtered is None or df_filtered.empty:
        return pd.DataFrame()

    kew_norm = str(kew_value or "").upper()
    df = df_filtered.copy()

    if "PROV" in kew_norm:
        key_col = "prov_key"
        label_cols = ("PROV_DISP", "KAB_DISP")
    else:
        key_col = "kab_key"
        label_cols = ("PROV_DISP", "KAB_DISP")

    if agg_total is None or agg_total.empty:
        df["Indeks_Final_0_100"] = df["Indeks_Dasar_0_100"]
    else:
        m = agg_total[["group_key", "Indeks_Final_Wilayah_0_100"]].copy()
        df = df.merge(m, left_on=key_col, right_on="group_key", how="left")
        df["Indeks_Final_0_100"] = df["Indeks_Final_Wilayah_0_100"].fillna(df["Indeks_Dasar_0_100"])
        df = df.drop(columns=[c for c in ["group_key","Indeks_Final_Wilayah_0_100"] if c in df.columns])

    base_cols = [label_cols[0], label_cols[1], "KEW_NORM", "_dataset"]
    if meta.get("nama_col") and meta["nama_col"] in df.columns:
        df["nm_perpustakaan"] = df[meta["nama_col"]].astype(str)
        base_cols.insert(2, "nm_perpustakaan")

    keep = base_cols + [
        "sub_koleksi","sub_sdm","sub_pelayanan","sub_pengelolaan",
        "dim_kepatuhan","dim_kinerja",
        "Indeks_Dasar_0_100",
        "Indeks_Final_0_100",
    ]
    keep = [c for c in keep if c in df.columns]

    out = df[keep].copy()
    out = out.rename(columns={label_cols[0]:"Provinsi", label_cols[1]:"Kab/Kota", "_dataset":"Jenis"})

    for c in ["sub_koleksi","sub_sdm","sub_pelayanan","sub_pengelolaan","dim_kepatuhan","dim_kinerja"]:
        if c in out.columns:
            out[c] = pd.to_numeric(out[c], errors="coerce").fillna(0.0).round(3)
    for c in ["Indeks_Dasar_0_100","Indeks_Final_0_100"]:
        if c in out.columns:
            out[c] = pd.to_numeric(out[c], errors="coerce").fillna(0.0).round(2)
    return out


# ============================================================
# 11) VERIF (kecukupan sampel)
# ============================================================

def build_verif_jenis(faktor_wilayah_jenis, kew_value):
    if faktor_wilayah_jenis is None or faktor_wilayah_jenis.empty:
        return pd.DataFrame()

    kew_norm = str(kew_value or "").upper()
    label_col = "Provinsi" if "PROV" in kew_norm else "Kab/Kota"

    out = faktor_wilayah_jenis.copy()
    keep = [c for c in [
        label_col, "Jenis",
        "pop_total_jenis", "target_total_33_88_jenis", "n_jenis",
        "coverage_jenis_%", "faktor_penyesuaian_jenis", "gap_target33_88_jenis"
    ] if c in out.columns]

    out = out[keep].copy()

    for c in ["pop_total_jenis", "target_total_33_88_jenis", "n_jenis", "gap_target33_88_jenis"]:
        if c in out.columns:
            out[c] = pd.to_numeric(out[c], errors="coerce").fillna(0).round(0).astype(int)
    if "coverage_jenis_%" in out.columns:
        out["coverage_jenis_%"] = pd.to_numeric(out["coverage_jenis_%"], errors="coerce").fillna(0.0).round(2)
    if "faktor_penyesuaian_jenis" in out.columns:
        out["faktor_penyesuaian_jenis"] = pd.to_numeric(out["faktor_penyesuaian_jenis"], errors="coerce").fillna(1.0).round(3)

    return out


# ============================================================
# 12) BELL CURVE — Indeks Dasar per Entitas (per Jenis) + Hover
# ============================================================

def _make_bell_curve_entitas(dfp, title, xcol="Indeks_Dasar_0_100", label_col="nm_perpustakaan", hover_cols=None, min_points=2):
    fig = go.Figure()
    fig.update_layout(
        title=title,
        xaxis_title="Skor (0–100)",
        yaxis_title="Kepadatan",
        hovermode="closest",
        margin=dict(l=40, r=20, t=60, b=40),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )

    if dfp is None or dfp.empty or xcol not in dfp.columns:
        fig.add_annotation(text="Tidak ada data untuk ditampilkan.", x=0.5, y=0.5, xref="paper", yref="paper", showarrow=False)
        fig.update_xaxes(range=[0, 100])
        fig.update_yaxes(rangemode="tozero")
        return fig

    d = dfp.dropna(subset=[xcol]).copy()
    if d.empty:
        fig.add_annotation(text="Tidak ada data untuk ditampilkan.", x=0.5, y=0.5, xref="paper", yref="paper", showarrow=False)
        fig.update_xaxes(range=[0, 100])
        fig.update_yaxes(rangemode="tozero")
        return fig

    x = pd.to_numeric(d[xcol], errors="coerce").astype(float)
    d = d.loc[x.notna()].copy()
    x = x.loc[x.notna()].values
    if len(x) < 1:
        fig.add_annotation(text="Tidak ada data untuk ditampilkan.", x=0.5, y=0.5, xref="paper", yref="paper", showarrow=False)
        fig.update_xaxes(range=[0, 100])
        fig.update_yaxes(rangemode="tozero")
        return fig

    hover_cols = hover_cols or []
    def _val(row, col):
        if col not in row.index:
            return ""
        v = row[col]
        return "" if pd.isna(v) else str(v)

    hover_text = []
    for _, row in d.iterrows():
        lines = []
        nm = _val(row, label_col) if (label_col and label_col in d.columns) else ""
        if nm:
            lines.append(f"<b>{nm}</b>")
        lines.append(f"{xcol}: {float(pd.to_numeric(row[xcol], errors='coerce')):.2f}")
        for hc in hover_cols:
            vv = _val(row, hc)
            if vv:
                lines.append(f"{hc}: {vv}")
        hover_text.append("<br>".join(lines))

    if len(x) < min_points:
        x_single = float(x[0])
        fig.add_trace(go.Scatter(
            x=[x_single], y=[0],
            mode="markers", showlegend=False,
            hovertext=[hover_text[0]] if hover_text else None,
            hoverinfo="text"
        ))
        fig.add_vline(x=x_single, line_width=1, line_dash="dash", annotation_text=f"Nilai: {x_single:.1f}", annotation_position="top")
        fig.update_xaxes(range=[0, 100])
        fig.update_yaxes(rangemode="tozero")
        return fig

    mu = float(np.mean(x))
    sigma = float(np.std(x, ddof=1)) if len(x) > 1 else 1.0
    sigma = max(sigma, 1e-3)

    xmin = max(0.0, float(np.min(x)) - 5.0)
    xmax = min(100.0, float(np.max(x)) + 5.0)
    xs = np.linspace(xmin, xmax, 250)
    pdf = (1.0 / (sigma * np.sqrt(2 * np.pi))) * np.exp(-0.5 * ((xs - mu) / sigma) ** 2)

    fig.add_trace(go.Scatter(x=xs, y=pdf, mode="lines", name="Kurva Normal (fit)"))
    fig.add_trace(go.Scatter(
        x=x, y=np.zeros_like(x),
        mode="markers", showlegend=False,
        hovertext=hover_text if hover_text else None,
        hoverinfo="text"
    ))

    q1, q2, q3 = np.percentile(x, [25, 50, 75])
    for xv, lab in [(q1, "Q1"), (q2, "Q2 (Median)"), (q3, "Q3"), (mu, "Mean")]:
        fig.add_vline(x=float(xv), line_width=1, line_dash="dash", annotation_text=f"{lab}: {xv:.1f}", annotation_position="top")

    fig.update_xaxes(range=[0, 100])
    fig.update_yaxes(rangemode="tozero")
    return fig


# ============================================================
# 13) KPI DASHBOARD (2 kartu: final + dasar)
# ============================================================

def _safe_first(df, col, default=0.0, where=None):
    if df is None or df.empty or col not in df.columns:
        return default
    sub = df
    if where is not None:
        sub = df.loc[where]
    if sub is None or sub.empty:
        return default
    return float(pd.to_numeric(sub[col], errors="coerce").fillna(default).iloc[0])

def build_kpi_markdown(summary_jenis):
    if summary_jenis is None or summary_jenis.empty:
        return ""
    final_all = _safe_first(summary_jenis, "Indeks_Final_Disesuaikan_0_100", 0.0, where=summary_jenis["Jenis"].astype(str).str.lower().eq("keseluruhan"))
    dasar_all = _safe_first(summary_jenis, "Indeks_Dasar_0_100", 0.0, where=summary_jenis["Jenis"].astype(str).str.lower().eq("keseluruhan"))

    def fmt(x, nd=2):
        return "NA" if pd.isna(x) else f"{x:.{nd}f}"

    return f"""
<div style="display:flex; gap:12px; flex-wrap:wrap;">
  <div style="border:1px solid #333; border-radius:10px; padding:10px 12px; min-width:260px;">
    <div style="opacity:0.8;">Indeks IPLM FINAL (Disesuaikan 33.88%)</div>
    <div style="font-size:26px; font-weight:700;">{fmt(final_all,2)}</div>
    <div style="opacity:0.7;">Skor absolut (untuk akuntabilitas)</div>
  </div>

  <div style="border:1px solid #333; border-radius:10px; padding:10px 12px; min-width:260px;">
    <div style="opacity:0.8;">Indeks Dasar (Tanpa Penyesuaian)</div>
    <div style="font-size:26px; font-weight:700;">{fmt(dasar_all,2)}</div>
    <div style="opacity:0.7;">Sebelum faktor kecukupan sampel</div>
  </div>
</div>
""".strip()


# ============================================================
# 14) LLM: Isi Interpretasi & Rekomendasi (TABEL) + WORD
# ============================================================

_HF_CLIENT = None

def get_llm_client():
    global _HF_CLIENT
    if _HF_CLIENT is not None:
        return _HF_CLIENT
    if (not HF_AVAILABLE) or (InferenceClient is None):
        _HF_CLIENT = None
        return None
    try:
        _HF_CLIENT = InferenceClient(model=LLM_MODEL_NAME, token=HF_TOKEN) if HF_TOKEN else InferenceClient(model=LLM_MODEL_NAME)
        return _HF_CLIENT
    except Exception:
        _HF_CLIENT = None
        return None

def _to_float(x, default=0.0):
    try:
        if x is None:
            return float(default)
        if isinstance(x, float) and math.isnan(x):
            return float(default)
        return float(x)
    except Exception:
        return float(default)

def summarize_jumlah_perpus_dari_agg_jenis(agg_jenis_full, wilayah_label, kew_value):
    """
    Ambil jumlah perpustakaan sumber data dari tabel agregat wilayah × jenis (gambar 2).
    Untuk filter 1 wilayah (kab/prov), agg_jenis_full biasanya 3 baris (sekolah/umum/khusus).
    Untuk nasional/semua wilayah, ini akan menjumlahkan seluruh wilayah per jenis.
    """
    if agg_jenis_full is None or agg_jenis_full.empty:
        return {"sekolah": 0, "umum": 0, "khusus": 0, "total": 0}

    a = agg_jenis_full.copy()
    if "Jenis" not in a.columns:
        return {"sekolah": 0, "umum": 0, "khusus": 0, "total": 0}

    a["Jenis"] = a["Jenis"].astype(str).str.lower().str.strip()
    if "Jumlah" in a.columns:
        a["Jumlah"] = pd.to_numeric(a["Jumlah"], errors="coerce").fillna(0).astype(int)
    else:
        a["Jumlah"] = 0

    out = {}
    for j in ["sekolah", "umum", "khusus"]:
        out[j] = int(a.loc[a["Jenis"].eq(j), "Jumlah"].sum())
    out["total"] = int(out["sekolah"] + out["umum"] + out["khusus"])
    return out

def build_interpretasi_table_values(agg_total, wilayah_label, target_ratio):
    """
    MENGAMBIL NILAI APA ADANYA (tanpa *100) dari hasil aplikasi (agg_total):
    - Kepatuhan = Rata2_dim_kepatuhan
    - Koleksi   = Rata2_sub_koleksi
    - Tenaga    = Rata2_sub_sdm
    - Kinerja   = Rata2_dim_kinerja
    - Pelayanan = Rata2_sub_pelayanan
    - Pengelolaan = Rata2_sub_pengelolaan
    - Nilai IPLM = Indeks_Final_Wilayah_0_100

    Jika agg_total > 1 baris (mis. nasional), diambil mean kolom-kolom tersebut.
    """
    if agg_total is None or agg_total.empty:
        base = {
            "kepatuhan": 0.0, "koleksi": 0.0, "tenaga": 0.0,
            "kinerja": 0.0, "pelayanan": 0.0, "pengelolaan": 0.0,
            "iplm": 0.0
        }
    else:
        a = agg_total.copy()
        cols_needed = [
            "Rata2_dim_kepatuhan",
            "Rata2_sub_koleksi",
            "Rata2_sub_sdm",
            "Rata2_dim_kinerja",
            "Rata2_sub_pelayanan",
            "Rata2_sub_pengelolaan",
            "Indeks_Final_Wilayah_0_100",
        ]
        for c in cols_needed:
            if c in a.columns:
                a[c] = pd.to_numeric(a[c], errors="coerce").fillna(0.0)
            else:
                a[c] = 0.0

        base = {
            "kepatuhan": float(a["Rata2_dim_kepatuhan"].mean()),
            "koleksi": float(a["Rata2_sub_koleksi"].mean()),
            "tenaga": float(a["Rata2_sub_sdm"].mean()),
            "kinerja": float(a["Rata2_dim_kinerja"].mean()),
            "pelayanan": float(a["Rata2_sub_pelayanan"].mean()),
            "pengelolaan": float(a["Rata2_sub_pengelolaan"].mean()),
            "iplm": float(a["Indeks_Final_Wilayah_0_100"].mean()),
        }

    # pembulatan display (nilai tetap "apa adanya", hanya format)
    # untuk sub/dim (0–1) biasanya 3 desimal; untuk IPLM (0–100) 2 desimal.
    base_disp = {
        "kepatuhan": round(_to_float(base["kepatuhan"]), 3),
        "koleksi": round(_to_float(base["koleksi"]), 3),
        "tenaga": round(_to_float(base["tenaga"]), 3),
        "kinerja": round(_to_float(base["kinerja"]), 3),
        "pelayanan": round(_to_float(base["pelayanan"]), 3),
        "pengelolaan": round(_to_float(base["pengelolaan"]), 3),
        "iplm": round(_to_float(base["iplm"]), 2),
    }

    rows = [
        {"No":"1",   "Dimensi":"Kepatuhan", "Nilai":base_disp["kepatuhan"], "SumberKolom":"Rata2_dim_kepatuhan"},
        {"No":"1.1", "Dimensi":"Variabel Koleksi", "Nilai":base_disp["koleksi"], "SumberKolom":"Rata2_sub_koleksi"},
        {"No":"1.2", "Dimensi":"Variabel Tenaga Perpustakaan", "Nilai":base_disp["tenaga"], "SumberKolom":"Rata2_sub_sdm"},
        {"No":"2",   "Dimensi":"Kinerja", "Nilai":base_disp["kinerja"], "SumberKolom":"Rata2_dim_kinerja"},
        {"No":"2.1", "Dimensi":"Variabel Pelayanan", "Nilai":base_disp["pelayanan"], "SumberKolom":"Rata2_sub_pelayanan"},
        {"No":"2.2", "Dimensi":"Variabel Penyelenggaraan/Pengelolaan", "Nilai":base_disp["pengelolaan"], "SumberKolom":"Rata2_sub_pengelolaan"},
        {"No":"4",   "Dimensi":"Nilai IPLM", "Nilai":base_disp["iplm"], "SumberKolom":"Indeks_Final_Wilayah_0_100"},
    ]

    header = {
        "judul": f"Interpretasi dan Rekomendasi IPLM — {wilayah_label}",
        "target_sampel": f"{target_ratio*100:.2f}%"
    }
    return header, rows

def _extract_json_from_llm(text):
    text = str(text or "").strip()

    text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*```$", "", text)

    try:
        return json.loads(text)
    except Exception:
        pass

    start = text.find("{")
    end = text.rfind("}")
    if start == -1 or end == -1 or end <= start:
        raise ValueError("Respons LLM tidak berisi JSON valid.")

    return json.loads(text[start:end + 1])

def fallback_fill_interpretasi_rekomendasi(rows, jumlah_perpus_by_jenis=None):
    """
    Fallback lokal saat LLM gagal/tidak tersedia.
    Tujuannya: kolom Interpretasi dan Rekomendasi tidak pernah kosong.
    """
    nilai = {
        str(r.get("Dimensi", "")).strip().lower(): _to_float(r.get("Nilai"), 0.0)
        for r in rows
    }

    kepatuhan = nilai.get("kepatuhan", 0.0)
    koleksi = nilai.get("variabel koleksi", 0.0)
    tenaga = nilai.get("variabel tenaga perpustakaan", 0.0)
    kinerja = nilai.get("kinerja", 0.0)
    pelayanan = nilai.get("variabel pelayanan", 0.0)
    pengelolaan = nilai.get("variabel penyelenggaraan/pengelolaan", 0.0)
    iplm = nilai.get("nilai iplm", 0.0)

    j = jumlah_perpus_by_jenis or {}
    total = int(j.get("total", 0) or 0)
    sekolah = int(j.get("sekolah", 0) or 0)
    umum = int(j.get("umum", 0) or 0)
    khusus = int(j.get("khusus", 0) or 0)

    sumber_data = (
        f"Sumber data mencakup {total} perpustakaan: sekolah {sekolah}, umum {umum}, khusus {khusus}."
        if total > 0
        else "Jumlah perpustakaan sumber data belum terbaca dari tabel agregat jenis."
    )

    def banding(a_label, a_val, b_label, b_val):
        gap = abs(a_val - b_val)
        if a_val > b_val:
            return f"{a_label} lebih besar daripada {b_label} dengan selisih {gap:.3f}"
        if a_val < b_val:
            return f"{a_label} lebih kecil daripada {b_label} dengan selisih {gap:.3f}"
        return f"{a_label} dan {b_label} berada pada nilai yang sama"

    template = {
        "kepatuhan": {
            "Interpretasi": (
                f"Nilai Kepatuhan sebesar {kepatuhan:.3f}. "
                f"Komponen pembentuknya menunjukkan bahwa {banding('koleksi', koleksi, 'tenaga perpustakaan', tenaga)}. "
                f"Angka ini menggambarkan gabungan kondisi koleksi dan kapasitas tenaga perpustakaan. {sumber_data}"
            ),
            "Rekomendasi": (
                "- Petakan satuan perpustakaan dengan nilai koleksi dan tenaga yang paling kecil.\n"
                "- Selaraskan pengembangan koleksi dengan penguatan kapasitas SDM.\n"
                "- Gunakan data per jenis perpustakaan untuk menentukan prioritas pendampingan."
            ),
        },
        "variabel koleksi": {
            "Interpretasi": (
                f"Nilai Variabel Koleksi sebesar {koleksi:.3f}. "
                f"Dibanding tenaga perpustakaan sebesar {tenaga:.3f}, posisinya menunjukkan bahwa "
                f"{banding('koleksi', koleksi, 'tenaga perpustakaan', tenaga)}. "
                f"Angka ini berkaitan dengan pengembangan, ketersediaan, penambahan, dan dukungan anggaran koleksi."
            ),
            "Rekomendasi": (
                "- Identifikasi kebutuhan judul dan eksemplar berdasarkan jenis perpustakaan.\n"
                "- Arahkan anggaran koleksi pada koleksi yang paling banyak dibutuhkan pemustaka.\n"
                "- Hubungkan pengadaan koleksi dengan data pemanfaatan layanan."
            ),
        },
        "variabel tenaga perpustakaan": {
            "Interpretasi": (
                f"Nilai Variabel Tenaga Perpustakaan sebesar {tenaga:.3f}. "
                f"Relasinya dengan koleksi menunjukkan bahwa {banding('tenaga perpustakaan', tenaga, 'koleksi', koleksi)}. "
                f"Angka ini menggambarkan kapasitas SDM, kualifikasi, fungsi profesional, pengembangan kompetensi, dan dukungan anggaran tenaga."
            ),
            "Rekomendasi": (
                "- Petakan kebutuhan tenaga berdasarkan kualifikasi dan fungsi layanan.\n"
                "- Prioritaskan pelatihan atau pengembangan kompetensi pada satuan yang nilainya masih tertinggal.\n"
                "- Sinkronkan kebutuhan SDM dengan rencana pengembangan koleksi dan layanan."
            ),
        },
        "kinerja": {
            "Interpretasi": (
                f"Nilai Kinerja sebesar {kinerja:.3f}. "
                f"Komponen pembentuknya menunjukkan bahwa {banding('pelayanan', pelayanan, 'penyelenggaraan/pengelolaan', pengelolaan)}. "
                f"Angka ini merepresentasikan gabungan aktivitas layanan dan tata kelola penyelenggaraan perpustakaan."
            ),
            "Rekomendasi": (
                "- Fokuskan tindak lanjut pada komponen yang nilainya lebih kecil antara pelayanan dan pengelolaan.\n"
                "- Selaraskan aktivitas layanan dengan kebijakan, kerja sama, variasi layanan, dan dukungan anggaran.\n"
                "- Pantau perubahan nilai pelayanan dan pengelolaan secara berkala."
            ),
        },
        "variabel pelayanan": {
            "Interpretasi": (
                f"Nilai Variabel Pelayanan sebesar {pelayanan:.3f}. "
                f"Dibanding penyelenggaraan/pengelolaan sebesar {pengelolaan:.3f}, polanya menunjukkan bahwa "
                f"{banding('pelayanan', pelayanan, 'penyelenggaraan/pengelolaan', pengelolaan)}. "
                f"Angka ini berkaitan dengan aktivitas layanan, budaya baca, pemustaka, fasilitas TIK, dan pemanfaatan koleksi."
            ),
            "Rekomendasi": (
                "- Perkuat kegiatan layanan yang langsung mendorong partisipasi dan pemanfaatan perpustakaan.\n"
                "- Integrasikan data layanan luring, daring, TIK, dan budaya baca dalam monitoring rutin.\n"
                "- Gunakan hasil per jenis perpustakaan untuk memilih sasaran intervensi layanan."
            ),
        },
        "variabel penyelenggaraan/pengelolaan": {
            "Interpretasi": (
                f"Nilai Variabel Penyelenggaraan/Pengelolaan sebesar {pengelolaan:.3f}. "
                f"Relasinya dengan pelayanan menunjukkan bahwa {banding('penyelenggaraan/pengelolaan', pengelolaan, 'pelayanan', pelayanan)}. "
                f"Angka ini menggambarkan tata kelola, kebijakan, kerja sama, variasi layanan, dan dukungan anggaran layanan."
            ),
            "Rekomendasi": (
                "- Lengkapi kebijakan dan prosedur layanan pada satuan perpustakaan sasaran.\n"
                "- Perkuat kerja sama dan variasi layanan yang mendukung kebutuhan pemustaka.\n"
                "- Tautkan perencanaan anggaran dengan program peningkatan pelayanan."
            ),
        },
        "nilai iplm": {
            "Interpretasi": (
                f"Nilai IPLM sebesar {iplm:.2f} pada skala 0-100. "
                f"Pembentuk utamanya menunjukkan bahwa {banding('kepatuhan', kepatuhan, 'kinerja', kinerja)}. "
                f"Nilai ini perlu dibaca bersama pola koleksi, tenaga perpustakaan, pelayanan, dan pengelolaan. {sumber_data}"
            ),
            "Rekomendasi": (
                "- Prioritaskan komponen yang nilainya lebih kecil antara kepatuhan dan kinerja.\n"
                "- Turunkan rencana tindak lanjut ke subvariabel koleksi, tenaga, pelayanan, dan pengelolaan.\n"
                "- Gunakan jumlah sumber data per jenis perpustakaan agar tindak lanjut lebih proporsional."
            ),
        },
    }

    out = []
    for r in rows:
        key = str(r.get("Dimensi", "")).strip().lower()
        isi = template.get(key)

        if isi is None:
            isi = {
                "Interpretasi": (
                    f"Nilai {r.get('Dimensi', '')} sebesar {r.get('Nilai', 0)} dibaca sebagai bagian dari struktur pembentuk IPLM."
                ),
                "Rekomendasi": (
                    "- Tinjau indikator penyusunnya.\n"
                    "- Cocokkan dengan data per jenis perpustakaan.\n"
                    "- Pantau perubahan pada periode berikutnya."
                ),
            }

        out.append({
            "No": str(r.get("No", "")),
            "Dimensi": str(r.get("Dimensi", "")),
            "Nilai": r.get("Nilai"),
            "Interpretasi": isi["Interpretasi"],
            "Rekomendasi": isi["Rekomendasi"],
        })

    return out

def llm_fill_interpretasi_rekomendasi(header, rows, wilayah_label, kew_label, jumlah_perpus_by_jenis):
    """
    LLM diminta mengisi kolom Interpretasi dan Rekomendasi dengan narasi yang NYAMBUNG ke angka:
    - Interpretasi: jelaskan apa arti angka untuk kondisi operasional perpustakaan (koleksi/sdm/pelayanan/pengelolaan),
      memakai relasi angka antardimensi (lebih besar/kecil, selisih, dominan, gap, konsistensi) TANPA label normatif.
    - Rekomendasi: 2–3 aksi teknis per baris yang langsung meng-address pola angka (misal dimensi lebih kecil → prioritas aktivitas),
      serta mengaitkan dengan volume data (jumlah perpustakaan per jenis) bila relevan.
    Output wajib JSON.
    """
    client = get_llm_client()
    if client is None or (not USE_LLM):
        return (
            fallback_fill_interpretasi_rekomendasi(rows, jumlah_perpus_by_jenis),
            "LLM tidak tersedia; Interpretasi & Rekomendasi diisi otomatis berbasis angka."
        )

    payload = {
        "wilayah": wilayah_label,
        "kewenangan": kew_label,
        "target_sampel_per_jenis": header["target_sampel"],
        "jumlah_perpustakaan_sumber_data": jumlah_perpus_by_jenis,
        "catatan_skala": (
            "Baris Kepatuhan/Koleksi/Tenaga/Kinerja/Pelayanan/Pengelolaan memakai nilai agregat 'apa adanya' "
            "(umumnya rentang 0–1 karena berasal dari sub/dim hasil normalisasi). "
            "Baris 'Nilai IPLM' memakai Indeks_Final_Wilayah_0_100 (rentang 0–100)."
        ),
        "baris": rows
    }

    system = (
        "Anda adalah analis kebijakan perpustakaan di Indonesia.\n"
        "Tugas: isi kolom Interpretasi dan Rekomendasi untuk setiap baris tabel.\n"
        "ATURAN WAJIB:\n"
        "1) Jangan mengubah nilai angka. Jangan menghitung ulang skor.\n"
        "2) Netral-deskriptif: dilarang memakai label normatif seperti baik/buruk, tinggi/sedang/rendah, memuaskan/kurang, optimal/tidak optimal.\n"
        "3) Interpretasi harus nyambung langsung ke angka dan relasinya antardimensi: gunakan istilah lebih besar/kecil, selisih, gap, dominan, konsisten/tidak konsisten, kontribusi, proporsi.\n"
        "4) Interpretasi juga harus menjelaskan kondisi riil berbasis dimensi:\n"
        "   - Koleksi: pengembangan, ketersediaan, pemanfaatan koleksi (sebagai fungsi layanan),\n"
        "   - Tenaga: kecukupan/kapasitas SDM dan pengembangan kompetensi,\n"
        "   - Pelayanan: aktivitas layanan dan pemanfaatan layanan,\n"
        "   - Pengelolaan: tata kelola, kebijakan, kolaborasi, dukungan anggaran layanan,\n"
        "   - Kepatuhan = gabungan koleksi+tenaga; Kinerja = gabungan pelayanan+pengelolaan.\n"
        "   Jelaskan tanpa menghakimi; fokus pada apa yang angka itu representasikan.\n"
        "5) Rekomendasi harus operasional dan spesifik (2–3 butir singkat) untuk tiap baris. Gunakan pola angka untuk menurunkan aksi.\n"
        "6) Output HARUS JSON valid saja (tanpa teks tambahan), dengan struktur persis.\n"
    )

    user = (
        "Kembalikan JSON:\n"
        "{\n"
        '  "rows": [\n'
        '    {"No":"...","Dimensi":"...","Nilai":..., "Interpretasi":"...","Rekomendasi":"..."}\n'
        "  ]\n"
        "}\n"
        "- Urutan dan jumlah baris harus sama.\n"
        "- 'Rekomendasi' boleh berupa bullet dengan tanda '-' dalam satu string.\n\n"
        f"INPUT:\n{json.dumps(payload, ensure_ascii=False)}"
    )

    try:
        resp = client.chat_completion(
            model=LLM_MODEL_NAME,
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
            max_tokens=1100,
            temperature=0.2,
            top_p=0.9,
        )
        text = resp.choices[0].message.content.strip()
        data = _extract_json_from_llm(text)
        rows_out = data.get("rows", [])
        if not isinstance(rows_out, list) or len(rows_out) != len(rows):
            raise ValueError("Format JSON rows tidak sesuai.")
        cleaned = []
        for i, r in enumerate(rows_out):
            cleaned.append({
                "No": str(r.get("No", rows[i]["No"])),
                "Dimensi": str(r.get("Dimensi", rows[i]["Dimensi"])),
                "Nilai": rows[i]["Nilai"],  # paksa nilai dari aplikasi
                "Interpretasi": str(r.get("Interpretasi","") or ""),
                "Rekomendasi": str(r.get("Rekomendasi","") or ""),
            })
        kosong = any((not r["Interpretasi"]) or (not r["Rekomendasi"]) for r in cleaned)
        if kosong:
            raise ValueError("LLM mengembalikan kolom kosong.")

        return cleaned, "LLM mengisi Interpretasi & Rekomendasi."
    except Exception as e:
        return (
            fallback_fill_interpretasi_rekomendasi(rows, jumlah_perpus_by_jenis),
            f"LLM error: {repr(e)}; Interpretasi & Rekomendasi diisi otomatis berbasis angka."
        )


def _set_cell_shading(cell, fill_hex="1F1F1F"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)

def _set_cell_text_color(cell, rgb_hex="FFFFFF"):
    for p in cell.paragraphs:
        for run in p.runs:
            rPr = run._r.get_or_add_rPr()
            color = OxmlElement("w:color")
            color.set(qn("w:val"), rgb_hex)
            rPr.append(color)

def _set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.append(tblPr)
    tblBorders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "8")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), "FFFFFF")
        tblBorders.append(elem)
    tblPr.append(tblBorders)

def generate_word_table_interpretasi(header, rows_filled, wilayah_label, jumlah_perpus_by_jenis):
    if (not DOCX_AVAILABLE) or (Document is None):
        return None

    doc = Document()

    # Title
    title = doc.add_paragraph()
    run = title.add_run(header["judul"])
    run.bold = True
    run.font.size = Pt(18)

    doc.add_paragraph(f"Target sampel per jenis: {header['target_sampel']}")

    # Table
    cols = ["No", "Dimensi", "Nilai", "Interpretasi", "Rekomendasi"]
    table = doc.add_table(rows=1, cols=len(cols))
    table.autofit = True
    _set_table_borders(table)

    hdr_cells = table.rows[0].cells
    for i, c in enumerate(cols):
        hdr_cells[i].text = c
        _set_cell_shading(hdr_cells[i], "1A1A1A")
        _set_cell_text_color(hdr_cells[i], "FFFFFF")
        for p in hdr_cells[i].paragraphs:
            for rr in p.runs:
                rr.bold = True

    for r in rows_filled:
        row_cells = table.add_row().cells
        row_cells[0].text = str(r.get("No",""))
        row_cells[1].text = str(r.get("Dimensi",""))

        # format nilai:
        # - sub/dim biasanya 0–1 → tampilkan 3 desimal
        # - IPLM 0–100 → tampilkan 2 desimal
        try:
            dim = str(r.get("Dimensi","")).strip().lower()
            val = _to_float(r.get("Nilai", 0.0), 0.0)
            if dim == "nilai iplm":
                row_cells[2].text = f"{val:.2f}"
            else:
                row_cells[2].text = f"{val:.3f}"
        except Exception:
            row_cells[2].text = str(r.get("Nilai",""))

        row_cells[3].text = str(r.get("Interpretasi","") or "")
        row_cells[4].text = str(r.get("Rekomendasi","") or "")

        for c in row_cells:
            _set_cell_shading(c, "262626")
            _set_cell_text_color(c, "FFFFFF")

    # ===== tambahan: deskripsi jumlah perpustakaan sumber data (gambar 2) =====
    doc.add_paragraph("")  # spacer
    j = jumlah_perpus_by_jenis or {"sekolah":0,"umum":0,"khusus":0,"total":0}
    p = doc.add_paragraph()
    p.add_run("Sumber data (jumlah perpustakaan pada tabel agregat wilayah × jenis): ").bold = True
    doc.add_paragraph(
        f"Perpustakaan sekolah = {int(j.get('sekolah',0))}, "
        f"perpustakaan umum = {int(j.get('umum',0))}, "
        f"perpustakaan khusus = {int(j.get('khusus',0))}, "
        f"total = {int(j.get('total',0))}."
    )

    outpath = tempfile.mktemp(suffix=".docx")
    doc.save(outpath)
    return outpath


# ============================================================
# 15) CORE RUN
# ============================================================

def _empty_outputs(msg="Data belum siap."):
    empty = pd.DataFrame()
    empty_fig = go.Figure()
    return (
        "",  # kpi_md
        empty, empty, empty, empty, empty,
        None, None, None, None, None,
        empty_fig, empty_fig, empty_fig,
        msg,  # msg
        "LLM belum tersedia.",  # status llm
        None  # word path
    )

def run_calc(prov_value, kab_value, kew_value, df_all, df_raw, pop_kab, pop_prov, pop_khusus, meta):
    try:
        if df_all is None or df_all.empty or df_raw is None or df_raw.empty:
            return _empty_outputs("Data belum ter-load. Pastikan file tersedia.")

        # Filter
        df = df_all.copy()
        if prov_value and prov_value != "(Semua)":
            df = df[df["PROV_DISP"] == prov_value]
        if kab_value and kab_value != "(Semua)":
            df = df[df["KAB_DISP"] == kab_value]
        if kew_value and kew_value != "(Semua)":
            df = df[df["KEW_NORM"] == kew_value]
        if df.empty:
            return _empty_outputs("Tidak ada data untuk filter ini.")

        kew_norm = kew_value if (kew_value and kew_value != "(Semua)") else "(Semua)"
        faktor_wilayah_jenis = build_faktor_wilayah_jenis(df, pop_kab, pop_prov, pop_khusus, kew_norm)
        agg_jenis_full = build_agg_wilayah_jenis(df, faktor_wilayah_jenis, kew_norm)
        agg_total = build_agg_wilayah_total_from_jenis(agg_jenis_full, faktor_wilayah_jenis, kew_norm)

        summary_jenis = build_summary_per_jenis(agg_jenis_full, agg_total)
        verif_total = build_verif_jenis(faktor_wilayah_jenis, kew_norm)
        detail_view = attach_final_to_detail(df, agg_total, meta, kew_norm)

        # agg_jenis view (UI hanya sampai indeks dasar)
        if agg_jenis_full is None or agg_jenis_full.empty:
            agg_jenis_view = agg_jenis_full
        else:
            kew_norm2 = str(kew_norm).upper()
            label_name = "Kab/Kota" if ("KAB" in kew_norm2 or "KOTA" in kew_norm2) else ("Provinsi" if "PROV" in kew_norm2 else "Kab/Kota")
            cols_upto = [
                "group_key",
                label_name,
                "Jenis",
                "Jumlah",
                "Rata2_sub_koleksi","Rata2_sub_sdm","Rata2_sub_pelayanan","Rata2_sub_pengelolaan",
                "Rata2_dim_kepatuhan","Rata2_dim_kinerja",
                "Indeks_Dasar_Agregat_0_100",
            ]
            cols_upto = [c for c in cols_upto if c in agg_jenis_full.columns]
            agg_jenis_view = agg_jenis_full[cols_upto].copy()

        # RAW download (hasil filter)
        raw = df_raw.copy()
        if prov_value and prov_value != "(Semua)":
            raw = raw[raw["PROV_DISP"] == prov_value]
        if kab_value and kab_value != "(Semua)":
            raw = raw[raw["KAB_DISP"] == kab_value]
        if kew_value and kew_value != "(Semua)":
            raw = raw[raw["KEW_NORM"] == kew_value]

        # Bell curve per jenis
        if detail_view is None or detail_view.empty:
            fig_umum = _make_bell_curve_entitas(pd.DataFrame(), "Bell Curve — Jenis: Umum")
            fig_sekolah = _make_bell_curve_entitas(pd.DataFrame(), "Bell Curve — Jenis: Sekolah")
            fig_khusus = _make_bell_curve_entitas(pd.DataFrame(), "Bell Curve — Jenis: Khusus")
        else:
            hover_cols = [hc for hc in ["Provinsi", "Kab/Kota", "Jenis"] if hc in detail_view.columns]

            def _fig(j):
                d = detail_view[detail_view["Jenis"].astype(str).str.lower() == j].copy()
                return _make_bell_curve_entitas(
                    d,
                    title=f"Bell Curve — Jenis: {j.title()} (Skor: Indeks_Dasar_0_100)",
                    xcol="Indeks_Dasar_0_100",
                    label_col=("nm_perpustakaan" if "nm_perpustakaan" in d.columns else "nm_perpustakaan"),
                    hover_cols=hover_cols,
                    min_points=2
                )

            fig_sekolah = _fig("sekolah")
            fig_umum = _fig("umum")
            fig_khusus = _fig("khusus")

        kpi_md = build_kpi_markdown(summary_jenis)

        # Export xlsx
        tmpdir = tempfile.mkdtemp()
        prov_slug = (_canon(prov_value or "SEMUA").upper() or "SEMUA")
        kab_slug = (_canon(kab_value or "SEMUA").upper() or "SEMUA")
        kew_slug = (_canon(kew_value or "SEMUA").upper() or "SEMUA")

        p_summary = str(Path(tmpdir) / f"IPLM_RingkasanJenisKeseluruhan_33_88_{prov_slug}_{kab_slug}_{kew_slug}.xlsx")
        p_total   = str(Path(tmpdir) / f"IPLM_AgregatWilayah_Keseluruhan_33_88_{prov_slug}_{kab_slug}_{kew_slug}.xlsx")
        p_raw     = str(Path(tmpdir) / f"IPLM_RAW_DATA_{prov_slug}_{kab_slug}_{kew_slug}.xlsx")
        p_detail  = str(Path(tmpdir) / f"IPLM_DetailEntitas_FinalMenempelWilayah_{prov_slug}_{kab_slug}_{kew_slug}.xlsx")
        p_verif   = str(Path(tmpdir) / f"IPLM_KecukupanSampel_33_88_{prov_slug}_{kab_slug}_{kew_slug}.xlsx")

        summary_jenis.to_excel(p_summary, index=False)
        agg_total.to_excel(p_total, index=False)
        raw.to_excel(p_raw, index=False)
        detail_view.to_excel(p_detail, index=False)
        verif_total.to_excel(p_verif, index=False)

        # ====== Word tabel interpretasi & rekomendasi ======
        wilayah_txt = kab_value if (kab_value and kab_value != "(Semua)") else (prov_value if (prov_value and prov_value != "(Semua)") else "Nasional/All")
        header, rows = build_interpretasi_table_values(agg_total, wilayah_txt, TARGET_RATIO)

        # jumlah perpustakaan sumber data (gambar 2)
        jumlah_perpus = summarize_jumlah_perpus_dari_agg_jenis(agg_jenis_full, wilayah_txt, kew_norm)

        rows_filled, llm_status = llm_fill_interpretasi_rekomendasi(
            header=header,
            rows=rows,
            wilayah_label=wilayah_txt,
            kew_label=(kew_value or "(Semua)"),
            jumlah_perpus_by_jenis=jumlah_perpus
        )
        word_path = generate_word_table_interpretasi(header, rows_filled, wilayah_txt, jumlah_perpus)

        msg = (
            f"Selesai (TARGET {TARGET_RATIO*100:.2f}%): raw={len(raw)} | entitas={len(detail_view)} | "
            f"wilayah(keseluruhan)={len(agg_total)} | jenis(agregat)={len(agg_jenis_full)}"
            + ("" if DOCX_AVAILABLE else " | python-docx tidak tersedia (Word OFF)")
        )

        return (
            kpi_md,
            summary_jenis, agg_total, agg_jenis_view, detail_view, verif_total,
            p_summary, p_total, p_raw, p_detail, p_verif,
            fig_umum, fig_sekolah, fig_khusus,
            msg,
            llm_status,
            (word_path if word_path else None)
        )

    except Exception as e:
        return _empty_outputs(f"Runtime error: {repr(e)}")


# ============================================================
# 16) UI (NO UPLOAD)
# ============================================================

def ui_load(force=False):
    df_all, df_raw, pop_kab, pop_prov, pop_khusus, meta, info = load_default_files(force=force)
    if df_all is None or (isinstance(df_all, pd.DataFrame) and df_all.empty):
        return (
            None, None, None, None, None, {}, info,
            gr.update(choices=["(Semua)"], value="(Semua)"),
            gr.update(choices=["(Semua)"], value="(Semua)"),
            gr.update(choices=["(Semua)"], value="(Semua)"),
        )

    prov_vals = df_all["PROV_DISP"].dropna().astype(str).tolist()
    prov_vals = [v for v in prov_vals if v and v.strip()]
    prov_choices = ["(Semua)"] + sorted(set(prov_vals))

    kab_choices = ["(Semua)"] + sorted([x for x in df_all["KAB_DISP"].dropna().unique().tolist() if x])
    kew_choices = ["(Semua)"] + sorted([x for x in df_all["KEW_NORM"].dropna().unique().tolist() if x])
    default_kew = "KAB/KOTA" if "KAB/KOTA" in kew_choices else ("PROVINSI" if "PROVINSI" in kew_choices else "(Semua)")

    return (
        df_all, df_raw, pop_kab, pop_prov, pop_khusus, meta, info,
        gr.update(choices=prov_choices, value="(Semua)"),
        gr.update(choices=kab_choices, value="(Semua)"),
        gr.update(choices=kew_choices, value=default_kew),
    )

def on_prov_change(prov_value):
    df_all, _, _, _, _, _, _ = load_default_files(force=False)
    if df_all is None or df_all.empty:
        return gr.update(choices=["(Semua)"], value="(Semua)")
    if prov_value is None or prov_value == "(Semua)":
        vals = df_all["KAB_DISP"].dropna().unique().tolist()
    else:
        vals = df_all.loc[df_all["PROV_DISP"] == prov_value, "KAB_DISP"].dropna().unique().tolist()
    vals = sorted([v for v in vals if v])
    return gr.update(choices=["(Semua)"] + vals, value="(Semua)")


with gr.Blocks() as demo:
    gr.Markdown(f"""
# IPLM 2025 — Final (Target Sampel 33.88% per Jenis) — TANPA Kinerja Relatif / Percentile
Mode NO UPLOAD (cache aktif). File dibaca dari repo/server:
- DATA_FILE   = {DATA_FILE}
- POP_KAB     = {POP_KAB}
- POP_PROV    = {POP_PROV}
- POP_KHUSUS  = {POP_KHUSUS}

TARGET RATIO (per jenis): {TARGET_RATIO*100:.2f}%

Dashboard KPI:
- Indeks IPLM FINAL (disesuaikan 33.88%)
- Indeks Dasar (tanpa penyesuaian)

UPDATE LLM + WORD:
- Tabel Word "Interpretasi & Rekomendasi" memakai NILAI APA ADANYA (tanpa dikali 100) untuk sub/dim.
- Baris "Nilai IPLM" memakai Indeks_Final_Wilayah_0_100 apa adanya.
- Di bawah tabel Word ditambahkan ringkasan jumlah perpustakaan sumber data (sekolah/umum/khusus/total) dari tabel agregat wilayah × jenis.
""")

    state_df = gr.State(None)
    state_raw = gr.State(None)
    state_pop_kab = gr.State(None)
    state_pop_prov = gr.State(None)
    state_pop_khusus = gr.State(None)
    state_meta = gr.State({})

    info_box = gr.Markdown()

    with gr.Row():
        dd_prov = gr.Dropdown(label="Provinsi", choices=["(Semua)"], value="(Semua)")
        dd_kab = gr.Dropdown(label="Kab/Kota", choices=["(Semua)"], value="(Semua)")
        dd_kew = gr.Dropdown(label="Kewenangan", choices=["(Semua)"], value="(Semua)")

    dd_prov.change(fn=on_prov_change, inputs=[dd_prov], outputs=dd_kab)

    run_btn = gr.Button("Jalankan Perhitungan")
    msg_out = gr.Markdown()

    kpi_out = gr.Markdown()

    gr.Markdown("## Ringkasan (Jenis + Keseluruhan)")
    out_summary = gr.DataFrame(interactive=False)

    gr.Markdown("## Agregat Wilayah (Keseluruhan) — FIX avg3")
    out_agg_total = gr.DataFrame(interactive=False)

    gr.Markdown("## Agregat Wilayah x Jenis — (ditampilkan sampai Indeks Dasar)")
    out_agg_jenis = gr.DataFrame(interactive=False)

    gr.Markdown("## Detail Entitas (Final menempel dari wilayah)")
    out_detail = gr.DataFrame(interactive=False)

    gr.Markdown("## Kecukupan Sampel 33.88%")
    out_verif = gr.DataFrame(interactive=False)

    gr.Markdown("## Bell Curve — Indeks Dasar per Entitas (per Jenis) + Nama Perpustakaan")
    gr.Markdown("### Perpustakaan Umum")
    bell_umum = gr.Plot(scale=1)

    gr.Markdown("### Perpustakaan Sekolah")
    bell_sekolah = gr.Plot(scale=1)

    gr.Markdown("### Perpustakaan Khusus")
    bell_khusus = gr.Plot(scale=1)

    gr.Markdown("## Status LLM (Isi Interpretasi & Rekomendasi)")
    llm_status_out = gr.Markdown()

    with gr.Row():
        dl_summary = gr.DownloadButton(label="Download Ringkasan (.xlsx)")
        dl_total = gr.DownloadButton(label="Download Agregat Wilayah (.xlsx)")
        dl_raw = gr.DownloadButton(label="Download Data Mentah (.xlsx)")
        dl_detail = gr.DownloadButton(label="Download Detail Entitas (.xlsx)")
        dl_verif = gr.DownloadButton(label="Download Kecukupan Sampel (.xlsx)")
        dl_word = gr.DownloadButton(label="Download Word: Interpretasi & Rekomendasi (.docx)" if DOCX_AVAILABLE else "Download Word (OFF)")

    run_btn.click(
        fn=run_calc,
        inputs=[dd_prov, dd_kab, dd_kew, state_df, state_raw, state_pop_kab, state_pop_prov, state_pop_khusus, state_meta],
        outputs=[
            kpi_out,
            out_summary, out_agg_total, out_agg_jenis, out_detail, out_verif,
            dl_summary, dl_total, dl_raw, dl_detail, dl_verif,
            bell_umum, bell_sekolah, bell_khusus,
            msg_out,
            llm_status_out,
            dl_word
        ]
    )

    demo.load(
        fn=lambda: ui_load(force=False),
        inputs=[],
        outputs=[state_df, state_raw, state_pop_kab, state_pop_prov, state_pop_khusus, state_meta, info_box, dd_prov, dd_kab, dd_kew]
    )

demo.launch()
